import csv
import io
from pathlib import Path

from sqlalchemy.exc import IntegrityError

from app.core.extensions import db
from app.models.user import User


REQUIRED_COLUMNS = {"nim", "nama", "password"}
ALLOWED_ROLES = {"mahasiswa", "fasil", "admin"}
MAX_ROWS = 2000


def _clean(value):
    return str(value).strip() if value is not None else ""


def _read_rows(filename: str, content: bytes) -> list[dict]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        text = content.decode("utf-8-sig")
        return [{_clean(key).lower(): _clean(value) for key, value in row.items()} for row in csv.DictReader(io.StringIO(text))]
    if suffix == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [_clean(value).lower() for value in rows[0]]
        return [dict(zip(headers, [_clean(value) for value in row])) for row in rows[1:] if any(row)]
    raise ValueError("Format file tidak didukung. Gunakan CSV atau XLSX.")


def import_users(filename: str, content: bytes) -> dict:
    rows = _read_rows(filename, content)
    if not rows:
        raise ValueError("File tidak berisi data pengguna.")
    if len(rows) > MAX_ROWS:
        raise ValueError(f"Maksimal {MAX_ROWS} baris dapat diimpor dalam satu batch.")

    columns = set(rows[0])
    missing = REQUIRED_COLUMNS - columns
    if missing:
        raise ValueError(f"Kolom wajib tidak ditemukan: {', '.join(sorted(missing))}.")

    errors = []
    normalized = []
    seen_nim = set()
    seen_email = set()
    for index, row in enumerate(rows, start=2):
        nim = _clean(row.get("nim"))
        nama = _clean(row.get("nama"))
        password = _clean(row.get("password"))
        role = _clean(row.get("role")) or "mahasiswa"
        email = _clean(row.get("email")) or None
        if not nim or not nama or not password:
            errors.append(f"Baris {index}: nim, nama, dan password wajib diisi.")
        if role not in ALLOWED_ROLES:
            errors.append(f"Baris {index}: role tidak valid.")
        if nim in seen_nim:
            errors.append(f"Baris {index}: NIM duplikat di dalam file.")
        if email and email in seen_email:
            errors.append(f"Baris {index}: email duplikat di dalam file.")
        seen_nim.add(nim)
        if email:
            seen_email.add(email)
        normalized.append({"nim": nim, "nama": nama, "password": password, "role": role, "email": email,
                           "asal": _clean(row.get("asal")) or None, "jekel": _clean(row.get("jekel")) or None,
                           "kamar_id": int(row["kamar_id"]) if _clean(row.get("kamar_id")).isdigit() else None})

    if errors:
        raise ValueError("Validasi file gagal: " + " ".join(errors[:20]))

    existing_nims = {value[0] for value in db.session.query(User.nim).filter(User.nim.in_(seen_nim)).all()}
    existing_emails = {value[0] for value in db.session.query(User.email).filter(User.email.in_(seen_email)).all()} if seen_email else set()
    if existing_nims or existing_emails:
        details = []
        if existing_nims:
            details.append("NIM sudah terdaftar: " + ", ".join(sorted(existing_nims)))
        if existing_emails:
            details.append("email sudah terdaftar: " + ", ".join(sorted(existing_emails)))
        raise ValueError("; ".join(details))

    users = []
    for data in normalized:
        password = data.pop("password")
        user = User(**data)
        user.set_password(password)
        users.append(user)
    try:
        db.session.add_all(users)
        db.session.commit()
    except IntegrityError as error:
        db.session.rollback()
        raise ValueError("Import dibatalkan karena ada data yang melanggar constraint database.") from error

    return {"created": len(users), "rows": len(rows)}